from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
from zipfile import BadZipFile

from django.db import transaction
from django.db.models import Q, Sum
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from payroll.models import Employee, EmployeeAlias, EmployeeRate, Payment, TimesheetEntry, normalize_person_name
from payroll.models import AccrualRecord, AuditLog


RATES_SOURCE = "rates"
TIMESHEET_SOURCE = "timesheet"
ITR_POSITIONS = {
    "Директор коммерческий",
    "Инженер",
    "Инженер ПТО",
    "Инженер по электромонтажу",
    "Ведущий инженер ПТО",
    "Координатор проекта",
    "Мастер",
    "Менеджер складского хозяйства",
    "Производитель работ",
    "Руководитель проекта",
}


@dataclass
class RateImportRow:
    row_number: int
    source_name: str
    organization: str
    project: str
    position: str
    employment_type: str
    hourly_rate: str
    daily_rate: str
    action: str
    employee_id: int | None
    personnel_number: str
    warnings: list[str]
    errors: list[str]

    @property
    def can_apply(self):
        return not self.errors

    def as_dict(self):
        return {
            "row_number": self.row_number,
            "source_name": self.source_name,
            "organization": self.organization,
            "project": self.project,
            "position": self.position,
            "employment_type": self.employment_type,
            "hourly_rate": self.hourly_rate,
            "daily_rate": self.daily_rate,
            "action": self.action,
            "employee_id": self.employee_id,
            "personnel_number": self.personnel_number,
            "warnings": self.warnings,
            "errors": self.errors,
        }


@dataclass
class TimesheetImportRow:
    row_number: int
    source_name: str
    days: int
    hours: str
    file_days: str
    file_hours: str
    daily_hours: dict[str, str]
    organization: str
    project: str
    position: str
    employment_type: str
    employment_type_label: str
    status: str
    status_label: str
    hourly_rate: str
    daily_rate: str
    provided_fields: list[str]
    action: str
    employee_id: int | None
    personnel_number: str
    has_rate: bool
    warnings: list[str]
    errors: list[str]

    def as_dict(self):
        return {
            "row_number": self.row_number,
            "source_name": self.source_name,
            "days": self.days,
            "hours": self.hours,
            "file_days": self.file_days,
            "file_hours": self.file_hours,
            "daily_hours": self.daily_hours,
            "organization": self.organization,
            "project": self.project,
            "position": self.position,
            "employment_type": self.employment_type,
            "employment_type_label": self.employment_type_label,
            "status": self.status,
            "status_label": self.status_label,
            "hourly_rate": self.hourly_rate,
            "daily_rate": self.daily_rate,
            "provided_fields": self.provided_fields,
            "action": self.action,
            "employee_id": self.employee_id,
            "personnel_number": self.personnel_number,
            "has_rate": self.has_rate,
            "warnings": self.warnings,
            "errors": self.errors,
        }


def cell_text(value):
    return "" if value is None else str(value).strip()


def normalize_employment_type(value):
    normalized = cell_text(value).casefold()
    if normalized == "итр":
        return Employee.EmploymentType.ITR
    if normalized in {"рабочий", "рабочий персонал"}:
        return Employee.EmploymentType.WORKER
    return None


def employment_type_for_position(position):
    return (
        Employee.EmploymentType.ITR
        if cell_text(position) in ITR_POSITIONS
        else Employee.EmploymentType.WORKER
    )


def normalize_organization(value):
    normalized = cell_text(value).casefold().replace("ё", "е")
    if not normalized:
        return ""
    if normalized.startswith("ип"):
        return Employee.Organization.INDIVIDUAL
    if normalized == "сз" or "самозанят" in normalized:
        return Employee.Organization.SELF_EMPLOYED
    if "подряд" in normalized:
        return Employee.Organization.CONTRACTOR
    if normalized == "штат" or "спектрум" in normalized or "сварта" in normalized:
        return Employee.Organization.SPECTRUM
    return None


def infer_employee_metadata(source_name, organization="", project="", position="", employment_type=None):
    name = normalize_person_name(source_name)
    inferred_type = employment_type or employment_type_for_position(position)
    inferred_organization = normalize_organization(organization)
    if not inferred_organization:
        if "подрядчик" in name:
            inferred_organization = Employee.Organization.CONTRACTOR
        elif "сз" in name:
            inferred_organization = Employee.Organization.SELF_EMPLOYED
        else:
            inferred_organization = Employee.Organization.SPECTRUM

    inferred_project = cell_text(project)
    if not inferred_project and "лобня" in name:
        inferred_project = "Лобня"

    return {
        "organization": inferred_organization,
        "project": inferred_project,
        "position": cell_text(position),
        "employment_type": inferred_type,
    }


def parse_decimal(value):
    if value in (None, ""):
        return ""
    if isinstance(value, Decimal):
        return str(value)
    text = str(value).replace("\xa0", "").replace(" ", "").replace(",", ".").strip()
    if not text:
        return ""
    try:
        return str(Decimal(text))
    except InvalidOperation:
        return None


def find_employee(source_name, source=None):
    normalized = normalize_person_name(source_name)
    aliases = EmployeeAlias.objects.select_related("employee").filter(normalized_source_name=normalized)
    if source:
        alias = aliases.filter(source=source).first()
        if alias is None:
            alias = aliases.first()
    else:
        alias = aliases.first()
    if alias:
        return alias.employee
    return Employee.objects.filter(normalized_full_name=normalized).first()


def parse_rates_workbook(file_obj):
    try:
        workbook = load_workbook(file_obj, data_only=True)
    except (BadZipFile, InvalidFileException, OSError) as exc:
        raise ValueError("Не удалось прочитать файл ставок. Загрузите корректный .xlsx файл.") from exc
    if "Ставки" not in workbook.sheetnames:
        raise ValueError("В файле не найден лист 'Ставки'.")

    sheet = workbook["Ставки"]
    rows = []
    seen_names = set()

    for row_number in range(4, sheet.max_row + 1):
        source_name = cell_text(sheet.cell(row=row_number, column=2).value)
        if not source_name:
            continue

        organization = cell_text(sheet.cell(row=row_number, column=3).value)
        project = cell_text(sheet.cell(row=row_number, column=4).value)
        position = cell_text(sheet.cell(row=row_number, column=5).value)
        employment_type_raw = cell_text(sheet.cell(row=row_number, column=6).value)
        normalized_organization = normalize_organization(organization or employment_type_raw)
        metadata = infer_employee_metadata(
            source_name,
            organization=normalized_organization or organization,
            project=project,
            position=position,
            employment_type=normalize_employment_type(employment_type_raw),
        )
        hourly_rate = parse_decimal(sheet.cell(row=row_number, column=7).value)
        daily_rate = parse_decimal(sheet.cell(row=row_number, column=8).value)

        warnings = []
        errors = []
        normalized_name = normalize_person_name(source_name)
        employee = find_employee(source_name, source=RATES_SOURCE)

        if organization and normalized_organization is None:
            errors.append(f"Неизвестная организация: {organization}.")

        if normalized_name in seen_names:
            errors.append("ФИО повторяется в файле ставок.")
        seen_names.add(normalized_name)

        if hourly_rate is None:
            errors.append("Некорректная ставка за час.")
            hourly_rate = ""
        if daily_rate is None:
            errors.append("Некорректная ставка за день.")
            daily_rate = ""
        if not hourly_rate:
            warnings.append("Ставка за час пустая.")
        if not daily_rate:
            warnings.append("Ставка за день пустая.")

        rows.append(
            RateImportRow(
                row_number=row_number,
                source_name=source_name,
                organization=metadata["organization"],
                project=metadata["project"],
                position=metadata["position"],
                employment_type=metadata["employment_type"],
                hourly_rate=hourly_rate,
                daily_rate=daily_rate,
                action="update" if employee else "create",
                employee_id=employee.id if employee else None,
                personnel_number=employee.personnel_number if employee else "",
                warnings=warnings,
                errors=errors,
            )
        )

    return rows


def parse_rates_upload(uploaded_file):
    return parse_rates_workbook(BytesIO(uploaded_file.read()))


def decimal_or_none(value):
    if value in (None, ""):
        return None
    return Decimal(str(value))


@transaction.atomic
def apply_rates_import(rows):
    created = 0
    updated = 0
    rates_updated = 0

    for row in rows:
        if row.get("errors"):
            continue

        employee = None
        if row.get("employee_id"):
            employee = Employee.objects.filter(pk=row["employee_id"]).first()
        if employee is None:
            employee = find_employee(row["source_name"], source=RATES_SOURCE)

        if employee is None:
            employee = Employee.objects.create(
                full_name=row["source_name"],
                organization=row["organization"],
                project=row["project"],
                position=row["position"],
                employment_type=row["employment_type"],
            )
            created += 1
        else:
            employee.organization = row["organization"]
            employee.project = row["project"]
            employee.position = row["position"]
            employee.employment_type = row["employment_type"]
            employee.save()
            updated += 1

        EmployeeAlias.objects.get_or_create(
            normalized_source_name=normalize_person_name(row["source_name"]),
            source=RATES_SOURCE,
            defaults={"employee": employee, "source_name": row["source_name"]},
        )

        EmployeeRate.objects.update_or_create(
            employee=employee,
            defaults={
                "hourly_rate": decimal_or_none(row["hourly_rate"]),
                "daily_rate": decimal_or_none(row["daily_rate"]),
            },
        )
        rates_updated += 1

    return {"created": created, "updated": updated, "rates_updated": rates_updated}


def parse_hour_value(value):
    parsed = parse_decimal(value)
    if parsed in (None, ""):
        return None
    return Decimal(parsed)


TIMESHEET_HEADERS = {
    "source_name": {"сотрудник", "фио"},
    "days": {"дни"},
    "hours": {"часы"},
    "organization": {"организация"},
    "project": {"проект"},
    "position": {"должность"},
    "hourly_rate": {"ставка ₽/час", "ставка руб/час", "ставка за час", "ставка час"},
    "daily_rate": {"ставка ₽/день", "ставка руб/день", "ставка за день", "ставка день"},
}


def normalize_header(value):
    return " ".join(cell_text(value).casefold().replace("ё", "е").split())


def date_header(value, period):
    if isinstance(value, (date, datetime)):
        parsed = value.date() if isinstance(value, datetime) else value
        return parsed.strftime("%d.%m") if (parsed.year, parsed.month) == (period.year, period.month) else ""

    match = re.match(r"^(\d{1,2})[./-](\d{1,2})(?:[./-](\d{2,4}))?(?:\s|$)", normalize_header(value))
    if not match:
        return ""
    day, month, year = match.groups()
    year = int(year) if year else period.year
    if year < 100:
        year += 2000
    try:
        parsed = date(year, int(month), int(day))
    except ValueError:
        return ""
    return parsed.strftime("%d.%m") if (parsed.year, parsed.month) == (period.year, period.month) else ""


def find_timesheet_layout(workbook, period):
    candidates = []
    for sheet in workbook.worksheets:
        for row_number in range(1, min(sheet.max_row, 20) + 1):
            columns = {}
            day_headers = {}
            for column in range(1, sheet.max_column + 1):
                value = sheet.cell(row=row_number, column=column).value
                header = normalize_header(value)
                for key, aliases in TIMESHEET_HEADERS.items():
                    if header in aliases and key not in columns:
                        columns[key] = column
                        break
                else:
                    if header.startswith("статус") and "status" not in columns:
                        columns["status"] = column
                    elif parsed_date := date_header(value, period):
                        day_headers[column] = parsed_date
            if {"source_name", "days", "hours"} <= columns.keys():
                candidates.append((len(day_headers), sheet, row_number, columns, day_headers))

    if not candidates:
        raise ValueError("Не найдены обязательные колонки 'Сотрудник', 'Дни' и 'Часы'.")
    best_day_count = max(item[0] for item in candidates)
    best = [item for item in candidates if item[0] == best_day_count]
    if len(best) != 1:
        locations = ", ".join(f"{item[1].title}, строка {item[2]}" for item in best)
        raise ValueError(f"Найдено несколько подходящих таблиц: {locations}.")
    _, sheet, row_number, columns, day_headers = best[0]
    return sheet, row_number, columns, day_headers


def parse_timesheet_workbook(file_obj, period):
    try:
        workbook = load_workbook(file_obj, data_only=True)
    except (BadZipFile, InvalidFileException, OSError) as exc:
        raise ValueError("Не удалось прочитать файл табеля. Загрузите корректный .xlsx файл.") from exc
    sheet, header_row, columns, day_headers = find_timesheet_layout(workbook, period)
    rows = []
    seen_names = set()

    for row_number in range(header_row + 1, sheet.max_row + 1):
        source_name = cell_text(sheet.cell(row=row_number, column=columns["source_name"]).value)
        normalized_name = normalize_person_name(source_name)
        is_section = any(
            merged.min_row <= row_number <= merged.max_row
            and merged.min_col <= columns["source_name"] <= merged.max_col
            and merged.max_col > merged.min_col
            for merged in sheet.merged_cells.ranges
        )
        if not source_name or normalized_name in {"итого", "уволенные"} or is_section:
            continue

        warnings = []
        errors = []
        employee = find_employee(source_name, source=TIMESHEET_SOURCE)

        if normalized_name in seen_names:
            errors.append("ФИО повторяется в табеле.")
        seen_names.add(normalized_name)

        daily_hours = {}
        for column, header in day_headers.items():
            parsed_hour = parse_hour_value(sheet.cell(row=row_number, column=column).value)
            if parsed_hour is None:
                continue
            if parsed_hour < 0:
                errors.append(f"Отрицательные часы в колонке {header}.")
                continue
            if parsed_hour == 0:
                continue
            daily_hours[header] = str(parsed_hour)

        file_days = cell_text(sheet.cell(row=row_number, column=columns["days"]).value)
        file_hours = cell_text(sheet.cell(row=row_number, column=columns["hours"]).value)
        file_days_decimal = parse_decimal(file_days)
        file_hours_decimal = parse_decimal(file_hours)
        if file_days_decimal is None:
            errors.append("Некорректное количество дней.")
            days = 0
        else:
            days_decimal = Decimal(file_days_decimal or "0")
            if days_decimal < 0 or days_decimal != days_decimal.to_integral_value():
                errors.append("Дни должны быть целым неотрицательным числом.")
                days = 0
            else:
                days = int(days_decimal)
        if file_hours_decimal is None:
            errors.append("Некорректное количество часов.")
            hours = Decimal("0")
        else:
            hours = Decimal(file_hours_decimal or "0")
            if hours < 0:
                errors.append("Часы не могут быть отрицательными.")

        provided_fields = [key for key in columns if key not in {"source_name", "days", "hours"}]
        organization_raw = cell_text(sheet.cell(row=row_number, column=columns["organization"]).value) if "organization" in columns else ""
        organization = normalize_organization(organization_raw)
        if organization_raw and organization is None:
            errors.append(f"Неизвестная организация: {organization_raw}.")
            organization = organization_raw
        project = cell_text(sheet.cell(row=row_number, column=columns["project"]).value) if "project" in columns else ""
        position = cell_text(sheet.cell(row=row_number, column=columns["position"]).value) if "position" in columns else ""
        employment_type = employment_type_for_position(position or (employee.position if employee else ""))
        status_raw = cell_text(sheet.cell(row=row_number, column=columns["status"]).value) if "status" in columns else ""
        status = ""
        if status_raw:
            status = Employee.Status.ARCHIVED if status_raw.casefold().startswith("уволен") else Employee.Status.ACTIVE
        hourly_rate = parse_decimal(sheet.cell(row=row_number, column=columns["hourly_rate"]).value) if "hourly_rate" in columns else ""
        daily_rate = parse_decimal(sheet.cell(row=row_number, column=columns["daily_rate"]).value) if "daily_rate" in columns else ""
        if hourly_rate is None:
            errors.append("Некорректная ставка за час.")
            hourly_rate = ""
        if daily_rate is None:
            errors.append("Некорректная ставка за день.")
            daily_rate = ""

        existing_rate = EmployeeRate.objects.filter(employee=employee).first() if employee else None
        effective_hourly_rate = hourly_rate if "hourly_rate" in columns else (existing_rate.hourly_rate if existing_rate else None)
        effective_daily_rate = daily_rate if "daily_rate" in columns else (existing_rate.daily_rate if existing_rate else None)
        has_rate = effective_hourly_rate not in (None, "") or effective_daily_rate not in (None, "")
        if not has_rate:
            warnings.append("Для сотрудника нет ставок.")

        rows.append(
            TimesheetImportRow(
                row_number=row_number,
                source_name=source_name,
                days=days,
                hours=str(hours),
                file_days=file_days,
                file_hours=file_hours,
                daily_hours=daily_hours,
                organization=organization or "",
                project=project,
                position=position,
                employment_type=employment_type,
                employment_type_label=Employee.EmploymentType(employment_type).label,
                status=status,
                status_label=Employee.Status(status).label if status else "",
                hourly_rate=hourly_rate,
                daily_rate=daily_rate,
                provided_fields=provided_fields,
                action="update" if employee else "create",
                employee_id=employee.id if employee else None,
                personnel_number=employee.personnel_number if employee else "",
                has_rate=has_rate,
                warnings=warnings,
                errors=errors,
            )
        )

    return rows


def parse_timesheet_upload(uploaded_file, period):
    return parse_timesheet_workbook(BytesIO(uploaded_file.read()), period)


@transaction.atomic
def apply_timesheet_import(period, rows):
    created_employees = 0
    updated_entries = 0

    for row in rows:
        if row.get("errors"):
            continue

        employee = None
        if row.get("employee_id"):
            employee = Employee.objects.filter(pk=row["employee_id"]).first()
        if employee is None:
            employee = find_employee(row["source_name"], source=TIMESHEET_SOURCE)

        if employee is None:
            employee = Employee.objects.create(
                full_name=row["source_name"],
                organization=row.get("organization") or Employee.Organization.SPECTRUM,
                project=row.get("project", ""),
                position=row.get("position", ""),
                employment_type=row.get("employment_type") or Employee.EmploymentType.WORKER,
                status=row.get("status") or Employee.Status.ACTIVE,
            )
            created_employees += 1
        else:
            provided_fields = set(row.get("provided_fields", []))
            changed_fields = []
            for field in ("organization", "project", "position", "status"):
                if field in provided_fields and row.get(field) and getattr(employee, field) != row[field]:
                    setattr(employee, field, row[field])
                    changed_fields.append(field)
            if "position" in provided_fields and row.get("position") and employee.employment_type != row["employment_type"]:
                employee.employment_type = row["employment_type"]
                changed_fields.append("employment_type")
            if changed_fields:
                employee.save(update_fields=[*changed_fields, "updated_at"])

        provided_fields = set(row.get("provided_fields", []))
        rate_fields = provided_fields & {"hourly_rate", "daily_rate"}
        if rate_fields:
            rate, _ = EmployeeRate.objects.get_or_create(employee=employee)
            for field in rate_fields:
                setattr(rate, field, decimal_or_none(row.get(field)))
            rate.save(update_fields=[*rate_fields, "updated_at"])

        EmployeeAlias.objects.get_or_create(
            normalized_source_name=normalize_person_name(row["source_name"]),
            source=TIMESHEET_SOURCE,
            defaults={"employee": employee, "source_name": row["source_name"]},
        )

        TimesheetEntry.objects.update_or_create(
            employee=employee,
            period=period,
            defaults={
                "days": row["days"],
                "hours": Decimal(str(row["hours"])),
                "daily_hours": row["daily_hours"],
                "source_name": row["source_name"],
            },
        )
        updated_entries += 1

    return {"created_employees": created_employees, "updated_entries": updated_entries}


def money(value):
    return value or Decimal("0")


ACCRUAL_TOTAL_KEYS = (
    "days",
    "hours",
    "accrued_days",
    "accrued_hours",
    "salary",
    "premium",
    "total",
    "total_with_ktu",
    "salary_1c",
    "daily_allowance_1c",
    "payments_total",
    "previous_plus",
    "previous_minus",
    "issued_total",
    "balance",
)


def calculate_accrual_totals(rows):
    totals = {key: Decimal("0") for key in ACCRUAL_TOTAL_KEYS}
    for row in rows:
        entry = row["entry"]
        record = row["record"]
        totals["days"] += Decimal(entry.days)
        totals["hours"] += money(entry.hours)
        totals["accrued_days"] += row["accrued_days"]
        totals["accrued_hours"] += row["accrued_hours"]
        totals["salary"] += money(record.salary)
        totals["premium"] += money(record.premium)
        totals["total"] += row["total"]
        totals["total_with_ktu"] += row["total_with_ktu"]
        totals["salary_1c"] += money(record.salary_1c)
        totals["daily_allowance_1c"] += money(record.daily_allowance_1c)
        totals["payments_total"] += row["payments_total"]
        totals["previous_plus"] += money(record.previous_plus)
        totals["previous_minus"] += money(record.previous_minus)
        totals["issued_total"] += row["issued_total"]
        totals["balance"] += row["balance"]
    return totals


def filter_accrual_rows(rows, filters):
    filtered_rows = []
    for row in rows:
        employee = row["employee"]
        if filters.get("only_zero") and not row["is_zero"]:
            continue
        if not filters.get("only_zero") and not filters.get("show_zero") and row["is_zero"]:
            continue
        if filters.get("fio") and normalize_person_name(filters["fio"]) not in employee.normalized_full_name:
            continue
        if filters.get("organization") and employee.organization != filters["organization"]:
            continue
        if filters.get("project") and employee.project != filters["project"]:
            continue
        if filters.get("position") and employee.position != filters["position"]:
            continue
        if filters.get("employment_type") and employee.employment_type != filters["employment_type"]:
            continue
        filtered_rows.append(row)
    return filtered_rows


def previous_period_for(period):
    year = period.year
    month = period.month - 1
    if month == 0:
        month = 12
        year -= 1
    return period.__class__.objects.filter(year=year, month=month).first()


def audit_change(instance, field_name, old_value, new_value, author):
    if str(old_value) == str(new_value):
        return
    AuditLog.objects.create(
        object_type=instance.__class__.__name__,
        object_id=instance.pk,
        object_repr=str(instance),
        field_name=field_name,
        old_value="" if old_value is None else str(old_value),
        new_value="" if new_value is None else str(new_value),
        author=author if getattr(author, "is_authenticated", False) else None,
    )


def audit_event(instance, message, author):
    AuditLog.objects.create(
        object_type=instance.__class__.__name__,
        object_id=instance.pk,
        object_repr=str(instance),
        field_name="event",
        old_value="",
        new_value=message,
        author=author if getattr(author, "is_authenticated", False) else None,
    )


def audit_deletion(object_type, object_id, object_repr, message, author):
    AuditLog.objects.create(
        object_type=object_type,
        object_id=object_id,
        object_repr=object_repr,
        field_name="event",
        old_value="",
        new_value=message,
        author=author if getattr(author, "is_authenticated", False) else None,
    )


def payment_created_message(payment):
    return (
        f"произвел выплату {payment.amount} руб. сотруднику "
        f"{payment.employee.full_name} за период {payment.period}; "
        f"дата выплаты {payment.paid_at:%d.%m.%Y}; инициатор: {payment.initiator}"
    )


def payment_updated_message(payment, changed_fields):
    field_labels = {
        "employee": "сотрудник",
        "period": "период",
        "amount": "сумма",
        "paid_at": "дата выплаты",
        "initiator": "инициатор",
        "comment": "комментарий",
    }
    changed = ", ".join(field_labels.get(field, field) for field in changed_fields)
    return (
        f"изменил выплату {payment.amount} руб. сотруднику "
        f"{payment.employee.full_name} за период {payment.period}"
        + (f"; изменены поля: {changed}" if changed else "")
    )


def audit_create(instance, fields, author):
    for field_name in fields:
        audit_change(instance, field_name, "", getattr(instance, field_name), author)


def audit_form_changes(instance, form, author):
    for field_name in form.changed_data:
        audit_change(
            instance,
            field_name,
            form.initial.get(field_name, ""),
            form.cleaned_data.get(field_name, ""),
            author,
        )


def sync_previous_balances(period):
    previous_period = previous_period_for(period)
    if previous_period is None:
        return 0

    previous_rows, _ = build_accrual_rows(previous_period)
    previous_balances = {row["employee"].id: row["balance"] for row in previous_rows}
    current_rows, _ = build_accrual_rows(period)
    updated = 0

    for row in current_rows:
        record = row["record"]
        if record.previous_balance_manual_override:
            continue

        previous_balance = money(previous_balances.get(row["employee"].id))
        next_plus = previous_balance if previous_balance > 0 else Decimal("0")
        next_minus = abs(previous_balance) if previous_balance < 0 else Decimal("0")

        if record.previous_plus != next_plus or record.previous_minus != next_minus:
            record.previous_plus = next_plus
            record.previous_minus = next_minus
            record.save(update_fields=["previous_plus", "previous_minus", "updated_at"])
            updated += 1

    return updated


def build_accrual_rows(period):
    entries = {
        entry.employee_id: entry
        for entry in TimesheetEntry.objects.select_related("employee", "period").filter(period=period)
    }
    employees = (
        Employee.objects.filter(Q(status=Employee.Status.ACTIVE) | Q(timesheet_entries__period=period))
        .distinct()
        .order_by("normalized_full_name", "id")
    )
    employee_ids = [employee.id for employee in employees]
    rates = {
        rate.employee_id: rate
        for rate in EmployeeRate.objects.filter(employee_id__in=employee_ids)
    }
    records = {
        record.employee_id: record
        for record in AccrualRecord.objects.filter(period=period, employee_id__in=employee_ids)
    }
    previous_period = previous_period_for(period)
    previous_salaries = (
        dict(
            AccrualRecord.objects.filter(period=previous_period, employee_id__in=employee_ids)
            .values_list("employee_id", "salary")
        )
        if previous_period
        else {}
    )
    payment_totals = {
        item["employee_id"]: item["total"] or Decimal("0")
        for item in Payment.objects.filter(period=period, employee_id__in=employee_ids)
        .values("employee_id")
        .annotate(total=Sum("amount"))
    }

    rows = []
    for employee in employees:
        entry = entries.get(employee.id) or TimesheetEntry(
            employee=employee,
            period=period,
            days=0,
            hours=Decimal("0"),
            daily_hours={},
        )
        rate = rates.get(employee.id)
        record = records.get(employee.id)
        if record is None:
            record = AccrualRecord.objects.create(
                employee=employee,
                period=period,
                salary=money(previous_salaries.get(employee.id)),
            )
            records[employee.id] = record

        daily_rate = money(rate.daily_rate) if rate else Decimal("0")
        hourly_rate = money(rate.hourly_rate) if rate else Decimal("0")
        accrued_days = Decimal(entry.days) * daily_rate
        accrued_hours = money(entry.hours) * hourly_rate
        premium = money(record.premium)
        is_zero = entry.days == 0 and money(entry.hours) == 0
        total_with_ktu = accrued_hours * money(record.ktu)
        total = money(record.salary) + premium if is_zero else total_with_ktu + accrued_days + premium
        payments_total = money(payment_totals.get(employee.id))
        issued_total = money(record.salary_1c) + money(record.daily_allowance_1c) + payments_total
        balance = (
            total
            - issued_total
            + money(record.previous_plus)
            - money(record.previous_minus)
        )

        row = {
            "employee": employee,
            "entry": entry,
            "rate": rate,
            "record": record,
            "daily_rate": daily_rate,
            "hourly_rate": hourly_rate,
            "accrued_days": accrued_days,
            "accrued_hours": accrued_hours,
            "total": total,
            "total_with_ktu": total_with_ktu,
            "payments_total": payments_total,
            "issued_total": issued_total,
            "balance": balance,
            "is_zero": is_zero,
        }
        rows.append(row)

    totals = calculate_accrual_totals(rows)
    return rows, totals
