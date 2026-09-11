from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from payroll.models import Payment
from payroll.services import build_accrual_rows, filter_accrual_rows, sync_previous_balances


HEADER_FILL = PatternFill("solid", fgColor="E9EEF5")
TOTAL_FILL = PatternFill("solid", fgColor="F6F8FA")


def workbook_response(workbook, filename):
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def format_sheet(sheet):
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL

    for column_cells in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        width = min(max(max_length + 2, 10), 42)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def export_accrual_workbook(period, filters=None):
    sync_previous_balances(period)
    rows, _ = build_accrual_rows(period)
    if filters:
        rows = filter_accrual_rows(rows, filters)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Начисление"
    sheet.append(["ФИО", "Остаток"])

    for row in rows:
        sheet.append([row["employee"].full_name, row["balance"]])

    for cell in sheet["B"]:
        if cell.row > 1:
            cell.number_format = "#,##0.00"

    format_sheet(sheet)
    return workbook


def export_payments_workbook(period):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Выплаты"
    sheet.append(
        [
            "Дата",
            "Таб. номер",
            "Сотрудник",
            "Сумма",
            "Инициатор",
            "Комментарий",
            "Автор изменения",
            "Дата изменения",
        ]
    )

    payments = (
        Payment.objects.select_related("employee", "updated_by")
        .filter(period=period)
        .order_by("paid_at", "employee__full_name")
    )
    total = 0
    for payment in payments:
        total += payment.amount
        sheet.append(
            [
                payment.paid_at,
                payment.employee.personnel_number,
                payment.employee.full_name,
                payment.amount,
                payment.initiator,
                payment.comment,
                payment.updated_by.username if payment.updated_by else "",
                payment.updated_at,
            ]
        )

    sheet.append(["", "Итого", "", total, "", "", "", ""])
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
        cell.fill = TOTAL_FILL

    for cell in sheet["A"]:
        if cell.row > 1:
            cell.number_format = "DD.MM.YYYY"
    for cell in sheet["D"]:
        if cell.row > 1:
            cell.number_format = "#,##0.00"
    for cell in sheet["H"]:
        if cell.row > 1:
            cell.number_format = "DD.MM.YYYY HH:MM"

    format_sheet(sheet)
    return workbook
