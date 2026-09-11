from decimal import Decimal
from io import BytesIO
from pathlib import Path

from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook

from core.models import AccountingPeriod
from payroll.exports import export_accrual_workbook
from payroll.models import AccrualRecord, AuditLog, Employee, EmployeeRate, Payment, TimesheetEntry
from payroll.services import (
    apply_timesheet_import,
    build_accrual_rows,
    filter_accrual_rows,
    find_timesheet_layout,
    parse_timesheet_workbook,
)
from users.models import User


class PayrollChangesTests(TestCase):
    def setUp(self):
        self.may = AccountingPeriod.objects.create(year=2026, month=5)
        self.june = AccountingPeriod.objects.create(year=2026, month=6)

    def employee(self, name="Иванов Иван", **kwargs):
        return Employee.objects.create(full_name=name, **kwargs)

    def workbook(self, headers, values, header_row=1, title="Любой лист"):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = title
        for column, header in enumerate(headers, 1):
            sheet.cell(header_row, column, header)
        for column, value in enumerate(values, 1):
            sheet.cell(header_row + 1, column, value)
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return output

    def test_header_based_import_and_metadata_apply(self):
        source = self.workbook(
            [
                "Примечание",
                "Часы",
                "Организация",
                "Сотрудник",
                "Дни",
                "Должность",
                "Статус на 30.08",
                "Ставка ₽/день",
                "01.06 Пн",
                "Ставка ₽/час",
            ],
            ["", 10, "ИП Русанов", "Петров Петр", 2, "Инженер ПТО", "Работает", 700, 5, 1000],
            header_row=3,
        )

        rows = parse_timesheet_workbook(source, self.june)
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual((row.days, row.hours), (2, "10"))
        self.assertEqual(row.daily_hours, {"01.06": "5"})
        self.assertEqual(row.organization, Employee.Organization.INDIVIDUAL)
        self.assertEqual(row.employment_type, Employee.EmploymentType.ITR)
        self.assertFalse(row.errors)

        apply_timesheet_import(self.june, [row.as_dict()])
        employee = Employee.objects.get(full_name="Петров Петр")
        rate = EmployeeRate.objects.get(employee=employee)
        self.assertEqual(employee.organization, Employee.Organization.INDIVIDUAL)
        self.assertEqual(employee.employment_type, Employee.EmploymentType.ITR)
        self.assertEqual((rate.hourly_rate, rate.daily_rate), (Decimal("1000"), Decimal("700")))
        self.assertEqual(TimesheetEntry.objects.get(employee=employee, period=self.june).hours, Decimal("10"))

    def test_blank_status_preserves_existing_employee_status(self):
        employee = self.employee("Архивный", status=Employee.Status.ARCHIVED)
        source = self.workbook(
            ["Сотрудник", "Дни", "Часы", "Статус"],
            [employee.full_name, 0, 0, ""],
        )
        row = parse_timesheet_workbook(source, self.june)[0]
        apply_timesheet_import(self.june, [row.as_dict()])
        employee.refresh_from_db()
        self.assertEqual(employee.status, Employee.Status.ARCHIVED)

    def test_import_normalizes_metadata_and_clears_provided_empty_rate(self):
        employee = self.employee(
            "Нормализуемый",
            organization=Employee.Organization.CONTRACTOR,
            status=Employee.Status.ACTIVE,
        )
        EmployeeRate.objects.create(employee=employee, hourly_rate=100, daily_rate=200)
        source = self.workbook(
            [
                "Сотрудник",
                "Дни",
                "Часы",
                "Организация",
                "Должность",
                "Статус сотрудника",
                "Ставка ₽/час",
            ],
            [employee.full_name, 0, 0, "Сварта", "Монтажник", "Уволен", ""],
        )
        row = parse_timesheet_workbook(source, self.june)[0]
        apply_timesheet_import(self.june, [row.as_dict()])

        employee.refresh_from_db()
        rate = EmployeeRate.objects.get(employee=employee)
        self.assertEqual(employee.organization, Employee.Organization.SPECTRUM)
        self.assertEqual(employee.employment_type, Employee.EmploymentType.WORKER)
        self.assertEqual(employee.status, Employee.Status.ARCHIVED)
        self.assertIsNone(rate.hourly_rate)
        self.assertEqual(rate.daily_rate, Decimal("200"))

    def test_layout_prefers_more_day_columns_and_rejects_tie(self):
        workbook = Workbook()
        first = workbook.active
        first.title = "Краткий"
        first.append(["Сотрудник", "Дни", "Часы", "01.06"])
        second = workbook.create_sheet("Полный")
        second.append(["Сотрудник", "Дни", "Часы", "01.06", "02.06"])
        sheet, *_ = find_timesheet_layout(workbook, self.june)
        self.assertEqual(sheet.title, "Полный")

        first.cell(1, 5, "02.06")
        with self.assertRaisesMessage(ValueError, "Найдено несколько подходящих таблиц"):
            find_timesheet_layout(workbook, self.june)

    def test_old_wide_format_uses_summary_columns(self):
        source = self.workbook(
            ["Сотрудник", "Дни", "Часы", "01.06", "02.06"],
            ["Сидоров Сидор", 7, 12, 5, 5],
            title="Табель",
        )
        row = parse_timesheet_workbook(source, self.june)[0]
        self.assertEqual((row.days, row.hours), (7, "12"))
        self.assertEqual(row.daily_hours, {"01.06": "5", "02.06": "5"})
        self.assertFalse(row.errors)

    def test_real_timesheet_has_expected_rows(self):
        path = Path(__file__).resolve().parent.parent / "tabel_2026-06-01_2026-06-30.xlsx"
        if not path.exists():
            self.skipTest("Контрольный XLSX не входит в production-образ.")
        rows = parse_timesheet_workbook(path, self.june)
        self.assertEqual(len(rows), 135)
        self.assertEqual(sum(row.days == 0 and Decimal(row.hours) == 0 for row in rows), 48)
        self.assertNotIn("УВОЛЕННЫЕ", {row.source_name for row in rows})
        self.assertNotIn("Итого", {row.source_name for row in rows})
        self.assertFalse([row for row in rows if row.errors])

    def test_calculation_zero_salary_and_salary_copy(self):
        worker = self.employee("Бета Рабочий")
        EmployeeRate.objects.create(employee=worker, hourly_rate=100, daily_rate=10)
        TimesheetEntry.objects.create(employee=worker, period=self.june, days=2, hours=3)
        AccrualRecord.objects.create(
            employee=worker,
            period=self.june,
            salary=999,
            premium=5,
            ktu=Decimal("1.5"),
            salary_1c=50,
            previous_plus=20,
            previous_minus=10,
        )
        zero = self.employee("Альфа Нулевой")
        may_record = AccrualRecord.objects.create(employee=zero, period=self.may, salary=500)

        rows, totals = build_accrual_rows(self.june)
        by_name = {row["employee"].full_name: row for row in rows}
        regular = by_name[worker.full_name]
        self.assertEqual(regular["total_with_ktu"], Decimal("450.0"))
        self.assertEqual(regular["total"], Decimal("475.0"))
        self.assertEqual(regular["balance"], Decimal("435.0"))
        self.assertFalse(regular["is_zero"])
        self.assertEqual(by_name[zero.full_name]["total"], Decimal("500"))
        self.assertTrue(by_name[zero.full_name]["is_zero"])
        self.assertEqual(totals["salary"], Decimal("1499"))

        may_record.salary = 700
        may_record.save()
        _, refreshed = build_accrual_rows(self.june)
        self.assertEqual(AccrualRecord.objects.get(employee=zero, period=self.june).salary, Decimal("500"))
        self.assertEqual(refreshed["total"], Decimal("975.0"))

    def test_filters_export_sorting_and_archived_history(self):
        beta = self.employee("бета", organization=Employee.Organization.CONTRACTOR, position="Сварщик")
        alpha = self.employee("Альфа", organization=Employee.Organization.SPECTRUM, position="Инженер")
        archived = self.employee("Ветеран", status=Employee.Status.ARCHIVED)
        self.employee("Уволен без табеля", status=Employee.Status.ARCHIVED)
        TimesheetEntry.objects.create(employee=beta, period=self.june, days=1, hours=1)
        TimesheetEntry.objects.create(employee=archived, period=self.june, days=1, hours=1)

        rows, _ = build_accrual_rows(self.june)
        self.assertEqual([row["employee"].full_name for row in rows], ["Альфа", "бета", "Ветеран"])
        self.assertEqual(
            [row["employee"].full_name for row in filter_accrual_rows(rows, {"fio": "БЕТ", "show_zero": True})],
            ["бета"],
        )
        self.assertNotIn(alpha, [row["employee"] for row in filter_accrual_rows(rows, {})])
        self.assertEqual(
            [row["employee"] for row in filter_accrual_rows(rows, {"only_zero": True})],
            [alpha],
        )
        workbook = export_accrual_workbook(self.june, {"fio": "вет", "show_zero": True})
        self.assertEqual(workbook.active["A2"].value, "Ветеран")
        self.assertIsNone(workbook.active["A3"].value)

        manager = User.objects.create_user(username="viewer", role=User.Role.MANAGER)
        self.client.force_login(manager)
        default_response = self.client.get(reverse("accrual_list"), {"period": self.june.pk})
        self.assertEqual(default_response.context["filter_options"]["positions"], ["Сварщик"])
        response = self.client.get(
            reverse("accrual_list"),
            {"period": self.june.pk, "fio": "альф", "show_zero": "1"},
        )
        self.assertEqual([row["employee"] for row in response.context["rows"]], [alpha])
        self.assertEqual(
            response.context["filter_options"]["organizations"],
            [Employee.Organization.CONTRACTOR, Employee.Organization.SPECTRUM],
        )
        response = self.client.get(
            reverse("accrual_list"),
            {"period": self.june.pk, "only_zero": "1"},
        )
        self.assertEqual([row["employee"] for row in response.context["rows"]], [alpha])
        self.assertEqual(response.context["filter_options"]["positions"], ["Инженер"])
        workbook = export_accrual_workbook(self.june, {"only_zero": True})
        self.assertEqual(workbook.active["A2"].value, "Альфа")
        self.assertIsNone(workbook.active["A3"].value)

    def test_position_filter_comes_from_employee_column_and_refreshes_after_edit(self):
        employee = self.employee("Редактируемый", position="Старая должность")
        admin = User.objects.create_user(username="position-admin", role=User.Role.ADMIN)
        self.client.force_login(admin)

        response = self.client.post(
            reverse("employee_update", args=[employee.pk]),
            {
                "full_name": employee.full_name,
                "organization": Employee.Organization.SPECTRUM,
                "project": "",
                "position": "Новая должность",
                "employment_type": Employee.EmploymentType.WORKER,
                "status": Employee.Status.ACTIVE,
            },
        )
        self.assertEqual(response.status_code, 302)
        response = self.client.get(
            reverse("accrual_list"),
            {"period": self.june.pk, "show_zero": "1"},
        )
        self.assertIn("Новая должность", response.context["filter_options"]["positions"])
        self.assertNotIn("Старая должность", response.context["filter_options"]["positions"])

    def test_payment_employee_field_is_searchable_datalist(self):
        employee = self.employee("Алексеев Алексей")
        manager = User.objects.create_user(username="payer", role=User.Role.MANAGER)
        self.client.force_login(manager)
        url = reverse("payment_create")

        response = self.client.get(url, {"period": self.june.pk})
        option = f"{employee.personnel_number} — {employee.full_name}"
        self.assertContains(response, 'list="employee-options"')
        self.assertContains(response, f'value="{option}"')

        response = self.client.post(
            url,
            {
                "employee": option,
                "period": self.june.pk,
                "amount": "1000",
                "paid_at": "2026-06-15",
                "comment": "",
            },
        )
        self.assertEqual(response.status_code, 302)
        payment = Payment.objects.get()
        self.assertEqual(payment.employee, employee)
        self.assertEqual(payment.initiator, manager.username)

    def test_only_admin_can_change_salary(self):
        employee = self.employee()
        record = AccrualRecord.objects.create(employee=employee, period=self.june, salary=100)
        manager = User.objects.create_user(username="manager", password="x", role=User.Role.MANAGER)
        admin = User.objects.create_user(username="admin", password="x", role=User.Role.ADMIN)
        url = reverse("accrual_edit", args=[record.pk])

        self.client.force_login(manager)
        self.client.post(url, {"salary": "900", "premium": "25", "ktu": "1"})
        record.refresh_from_db()
        self.assertEqual(record.salary, Decimal("100"))
        self.assertEqual(record.premium, Decimal("25"))

        self.client.force_login(admin)
        response = self.client.post(
            url,
            {
                "salary": "900",
                "premium": "0",
                "ktu": "1",
                "salary_1c": "0",
                "daily_allowance_1c": "0",
                "previous_plus": "0",
                "previous_minus": "0",
            },
        )
        self.assertEqual(response.status_code, 302, response.context["form"].errors if response.context else "")
        record.refresh_from_db()
        self.assertEqual(record.salary, Decimal("900"))
        self.assertTrue(AuditLog.objects.filter(object_id=record.pk, field_name="salary").exists())
